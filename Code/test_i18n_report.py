# -*- coding: utf-8 -*-
"""真实生成一份中文报告，验证模板确实换了语言。

不动用户的项目，用一个临时 payload。生成的文件放在 report_cache 里，测完删掉。
"""
import sys
import time
from pathlib import Path

sys.path.insert(0, r"<仓库根>\Code")
import i18n
import project_manager as pm

ok = True


def chk(cond, label, extra=""):
    global ok
    if not cond:
        ok = False
    print("  [%s] %s%s" % ("OK" if cond else "FAIL", label,
                           (" -> " + str(extra)) if extra else ""))


PAYLOAD = {
    "project_name": "双语测试",
    "image_mode": "full",
    "points": [
        # normalize_report_points 要的是 lat/lng/time 这些**原始字段**，
        # 格式化文本一律由服务端生成（它不信前端）—— 一开始我传了
        # lat_text/time_text，全被过滤掉，报"没有有效的标记点"。
        {"name": "Point A", "lat": 23.054897, "lng": 114.108218,
         "place": "Guangdong, Huizhou", "video": "DJI_0001_V.MP4",
         "time": 12.3, "zoom": 17, "basemap_name": "Street"},
    ],
}


def text_of(path):
    """把 docx 里的文字全读出来（用于断言语言）。"""
    from docx import Document
    d = Document(str(path))
    parts = []
    for p in d.paragraphs:
        parts.append(p.text)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


made = []
try:
    for lang, must_have, must_not in (
        ("zh", ["标记点报告", "生成时间", "目录", "汇总", "视频文件", "视频时刻"],
         ["Marker Report", "Generated:", "Contents", "Summary", "Video file"]),
        ("en", ["Marker Report", "Generated:", "Contents", "Summary", "Video file"],
         ["标记点报告", "生成时间", "视频文件"]),
    ):
        payload = dict(PAYLOAD)
        payload["lang"] = lang
        info = pm.build_docx_report(payload)
        made.append(pm.REPORT_DIR / info["file"])
        body = text_of(pm.REPORT_DIR / info["file"])
        print("=== %s 报告: %s ===" % (lang, info["file"]))
        print("  file=%s" % info["file"])
        print("  download_name=%s" % info["download_name"])
        for s in must_have:
            chk(s in body, "含 %r" % s)
        for s in must_not:
            chk(s not in body, "不含 %r" % s)
        print()
finally:
    for p in made:
        try:
            p.unlink()
        except OSError:
            pass
    print("  临时报告已删除: %s" % all(not p.exists() for p in made))

print()
print("结果:", "全部通过" if ok else "有失败项")
raise SystemExit(0 if ok else 1)
